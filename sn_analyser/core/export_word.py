# -*- coding: utf-8 -*-
"""
core/export_word.py
-----------------------
Generation du rapport Word (mise en forme narrative + synthese + tableau
de resultats) a partir des resultats d analyse.
"""

import datetime


def _verifier_dependance():
    """Message explicite si la bibliotheque d export n est pas installee
    (elle n est pas fournie par defaut avec QGIS)."""
    try:
        __import__("docx")
    except ImportError:
        raise RuntimeError(
            "La bibliotheque Python 'python-docx' n'est pas installee. Sous Windows, ouvrez l'OSGeo4W Shell "
            "et tapez : python -m pip install python-docx  (puis redemarrez QGIS)."
        )


def generer_rapport_word(nom, annee_debut, annee_fin, resultats, commentaire, chemin_sortie):
    """Assemble un rapport Word a partir des resultats d analyse.

    :param nom: nom de la zone d etude
    :param annee_debut: annee de debut
    :param annee_fin: annee de fin
    :param resultats: dict des resultats (cles voir core.utils.INDICATOR_LABELS)
    :param commentaire: commentaire de synthese (core.commentaire_ia)
    :param chemin_sortie: chemin du fichier .docx a produire
    """
    _verifier_dependance()
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from .utils import INDICATOR_LABELS

    doc = Document()

    titre = doc.add_heading("SN ANALYSER", level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_titre = doc.add_paragraph("Rapport d analyse territoriale")
    sous_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_titre.runs[0].font.size = Pt(14)
    sous_titre.runs[0].font.color.rgb = RGBColor(0x2E, 0x8B, 0x57)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Zone d etude : ").bold = True
    p.add_run(nom)
    p2 = doc.add_paragraph()
    p2.add_run("Date analysee : ").bold = True
    p2.add_run(str(annee_fin) if annee_debut == annee_fin else str(annee_debut) + " - " + str(annee_fin))
    p3 = doc.add_paragraph()
    p3.add_run("Date de generation : ").bold = True
    p3.add_run(datetime.datetime.now().strftime("%d/%m/%Y a %H:%M"))

    doc.add_heading("Synthese", level=1)
    doc.add_paragraph(commentaire)

    doc.add_heading("Resultats detailles", level=1)
    lignes_disponibles = [(lbl, resultats[cle]) for cle, lbl in INDICATOR_LABELS.items() if cle in resultats]

    if lignes_disponibles:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Indicateur"
        hdr[1].text = "Valeur"
        for libelle, valeur in lignes_disponibles:
            row = table.add_row().cells
            row[0].text = libelle
            row[1].text = ("%.2f" % valeur) if isinstance(valeur, float) else str(valeur)
    else:
        doc.add_paragraph("Aucun resultat detaille disponible.")

    doc.add_paragraph()
    footer = doc.add_paragraph("Genere automatiquement par le plugin QGIS SN Analyser (Geo Senegal 2026).")
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(chemin_sortie)
    return chemin_sortie
